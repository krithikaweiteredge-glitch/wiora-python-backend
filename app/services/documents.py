"""Extract text from an uploaded file (blueprint §13). PDF via PyMuPDF, DOCX via
python-docx, XLSX via openpyxl, plain text directly, and images via Gemini vision
OCR (reusing the configured vision model)."""
from __future__ import annotations

import base64
import io

MAX_CHARS = 20000


def _ocr_image(mimetype: str, data: bytes) -> str:
    """OCR an image by asking the vision model to transcribe its text."""
    try:
        from ..ai.service import ai_service

        uri = f"data:{mimetype};base64,{base64.b64encode(data).decode()}"
        text = ai_service.generate_vision(
            "You are an OCR engine.",
            "Transcribe ALL text visible in this image exactly. Return only the text, "
            "no commentary. If there is no text, reply with an empty string.",
            uri,
        )
        return (text or "").strip()[:MAX_CHARS]
    except Exception:
        return ""


def _extract_xlsx(data: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append("\t".join(cells))
    return "\n".join(lines)[:MAX_CHARS]


def extract_text(filename: str, mimetype: str, data: bytes) -> str:
    name = (filename or "").lower()
    mt = (mimetype or "").lower()
    try:
        if mt == "application/pdf" or name.endswith(".pdf"):
            import fitz  # PyMuPDF

            doc = fitz.open(stream=data, filetype="pdf")
            return "\n".join(page.get_text() for page in doc)[:MAX_CHARS]
        if name.endswith(".docx") or "word" in mt or "officedocument.wordprocessing" in mt:
            import docx

            d = docx.Document(io.BytesIO(data))
            return "\n".join(p.text for p in d.paragraphs)[:MAX_CHARS]
        if name.endswith((".xlsx", ".xlsm")) or "spreadsheet" in mt or "excel" in mt:
            return _extract_xlsx(data)
        if mt.startswith("text/") or name.endswith((".txt", ".md", ".csv", ".json")):
            return data.decode("utf-8", "ignore")[:MAX_CHARS]
        if mt.startswith("image/"):
            return _ocr_image(mt, data)  # Gemini vision OCR
    except Exception:
        return ""
    # Unknown type — best-effort decode.
    try:
        return data.decode("utf-8", "ignore")[:MAX_CHARS]
    except Exception:
        return ""


def chunk_text(text: str, size: int = 800, overlap: int = 120) -> list[str]:
    """Split text into overlapping chunks for embedding/retrieval."""
    text = " ".join((text or "").split())
    if not text:
        return []
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = end - overlap
    return chunks[:50]  # cap chunks per document
